#!/usr/bin/env python3
"""Command PDF — local PDF reading, conversion, editing, and signing."""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PySide6.QtCore import QEvent, QPoint, QPointF, Qt
from PySide6.QtGui import QAction, QColor, QIcon, QImage, QPainter, QPen
from PySide6.QtPdf import QPdfDocument
from PySide6.QtPdfWidgets import QPdfView
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QDoubleSpinBox, QFileDialog, QFormLayout,
    QFrame, QHBoxLayout, QLabel, QLineEdit, QListWidget, QMainWindow, QMessageBox,
    QPlainTextEdit, QPushButton, QSpinBox, QSplitter, QToolBar, QVBoxLayout, QWidget,
)

APP_DIR = Path(__file__).resolve().parent
ICON = APP_DIR / "assets" / "command-pdf.png"


def require_fitz(parent):
    try:
        import fitz
        return fitz
    except ImportError:
        QMessageBox.critical(parent, "Editing unavailable", "PyMuPDF is missing. Run setup.sh in the Command PDF folder.")
        return None


def output_path(parent, title, suggested):
    path, _ = QFileDialog.getSaveFileName(parent, title, str(suggested), "PDF documents (*.pdf)")
    return Path(path) if path else None


class SignaturePad(QWidget):
    def __init__(self):
        super().__init__()
        self.setMinimumSize(560, 210)
        self.image = QImage(1120, 420, QImage.Format_ARGB32)
        self.image.fill(Qt.transparent)
        self.last = None

    def clear(self):
        self.image.fill(Qt.transparent)
        self.update()

    def mousePressEvent(self, event):
        if event.button() == Qt.LeftButton:
            self.last = event.position()

    def mouseMoveEvent(self, event):
        if self.last is None or not event.buttons() & Qt.LeftButton:
            return
        scale_x, scale_y = self.image.width() / self.width(), self.image.height() / self.height()
        painter = QPainter(self.image)
        painter.setRenderHint(QPainter.Antialiasing)
        painter.setPen(QPen(QColor("#0b1320"), 7, Qt.SolidLine, Qt.RoundCap, Qt.RoundJoin))
        painter.drawLine(QPointF(self.last.x() * scale_x, self.last.y() * scale_y),
                         QPointF(event.position().x() * scale_x, event.position().y() * scale_y))
        painter.end()
        self.last = event.position()
        self.update()

    def mouseReleaseEvent(self, _event):
        self.last = None

    def paintEvent(self, _event):
        painter = QPainter(self)
        painter.fillRect(self.rect(), QColor("#ffffff"))
        painter.setPen(QPen(QColor("#d3dfeb"), 1))
        painter.drawRect(self.rect().adjusted(0, 0, -1, -1))
        painter.drawImage(self.rect(), self.image)
        painter.setPen(QPen(QColor("#c5d2df"), 1, Qt.DashLine))
        painter.drawLine(24, self.height() - 38, self.width() - 24, self.height() - 38)


class SignatureDialog(QDialog):
    def __init__(self, page_count, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Add visible signature")
        self.resize(620, 390)
        layout = QVBoxLayout(self)
        note = QLabel("Draw your signature below. This creates a visible signature stamp; it is not a certificate-based digital signature.")
        note.setWordWrap(True)
        layout.addWidget(note)
        self.pad = SignaturePad()
        layout.addWidget(self.pad)
        form = QFormLayout()
        self.page = QSpinBox(); self.page.setRange(1, max(1, page_count)); self.page.setValue(max(1, page_count))
        self.position = QComboBox(); self.position.addItems(["Bottom right", "Bottom left", "Top right", "Top left", "Centre"])
        self.width = QSpinBox(); self.width.setRange(60, 300); self.width.setValue(150); self.width.setSuffix(" pt")
        form.addRow("Page", self.page); form.addRow("Position", self.position); form.addRow("Width", self.width)
        layout.addLayout(form)
        row = QHBoxLayout(); clear = QPushButton("Clear"); clear.clicked.connect(self.pad.clear); row.addWidget(clear); row.addStretch()
        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok)
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); row.addWidget(buttons); layout.addLayout(row)


class TextDialog(QDialog):
    def __init__(self, page_count, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Add text box")
        self.resize(520, 470)
        layout = QVBoxLayout(self); form = QFormLayout()
        self.page = QSpinBox(); self.page.setRange(1, max(1, page_count))
        self.text = QPlainTextEdit(); self.text.setPlaceholderText("Type the text for this box…"); self.text.setMinimumHeight(120)
        self.position = QComboBox(); self.position.addItems(["Top left", "Top centre", "Bottom left", "Bottom centre", "Centre"])
        self.size = QSpinBox(); self.size.setRange(6, 72); self.size.setValue(12); self.size.setSuffix(" pt")
        self.width = QDoubleSpinBox(); self.width.setRange(72, 500); self.width.setValue(280); self.width.setSuffix(" pt")
        self.height = QDoubleSpinBox(); self.height.setRange(36, 500); self.height.setValue(110); self.height.setSuffix(" pt")
        self.border = QCheckBox("Draw a border around the text box"); self.border.setChecked(True)
        self.background = QCheckBox("Use a white background"); self.background.setChecked(True)
        form.addRow("Page", self.page); form.addRow("Text", self.text); form.addRow("Position", self.position)
        form.addRow("Box width", self.width); form.addRow("Box height", self.height); form.addRow("Font size", self.size)
        form.addRow("", self.border); form.addRow("", self.background)
        layout.addLayout(form)
        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok)
        buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); layout.addWidget(buttons)


class CommandPDF(QMainWindow):
    def __init__(self, initial=None):
        super().__init__()
        self.setWindowTitle("Command PDF")
        self.setWindowIcon(QIcon(str(ICON)))
        self.resize(1280, 820)
        self.setAcceptDrops(True)
        self.current = None
        self.document = QPdfDocument(self)
        self.document.pageCountChanged.connect(self.document_ready)
        self.viewer = QPdfView()
        self.viewer.setDocument(self.document)
        self.viewer.setPageMode(QPdfView.PageMode.MultiPage)
        self.viewer.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        self.viewer.viewport().installEventFilter(self)
        self.viewer.grabGesture(Qt.PinchGesture)
        self.pages = QListWidget(); self.pages.setMaximumWidth(170); self.pages.currentRowChanged.connect(self.go_page)
        self.status = QLabel("Open a PDF to begin")
        self.status.setObjectName("status")
        self.build_ui()
        if initial:
            self.open_pdf(Path(initial))

    def eventFilter(self, watched, event):
        if watched is self.viewer.viewport() and event.type() == QEvent.Wheel:
            pixels, angles = event.pixelDelta(), event.angleDelta()
            modifiers = event.modifiers()
            if modifiers & Qt.ControlModifier:
                delta = pixels.y() or angles.y()
                if delta:
                    self.zoom_by(pow(1.0018, delta))
                    event.accept()
                    return True
            if modifiers & Qt.ShiftModifier and not pixels.x():
                delta = pixels.y() or angles.y() / 2
                bar = self.viewer.horizontalScrollBar()
                bar.setValue(bar.value() - int(delta))
                event.accept()
                return True
            if pixels.x():
                horizontal = self.viewer.horizontalScrollBar()
                vertical = self.viewer.verticalScrollBar()
                horizontal.setValue(horizontal.value() - pixels.x())
                if pixels.y():
                    vertical.setValue(vertical.value() - pixels.y())
                event.accept()
                return True
        if watched is self.viewer.viewport() and event.type() == QEvent.NativeGesture:
            gesture_type = event.gestureType()
            zoom_type = getattr(Qt, "ZoomNativeGesture", Qt.NativeGestureType.ZoomNativeGesture)
            if gesture_type == zoom_type:
                self.zoom_by(1.0 + event.value())
                event.accept()
                return True
        return super().eventFilter(watched, event)

    def event(self, event):
        if event.type() == QEvent.Gesture:
            pinch = event.gesture(Qt.PinchGesture)
            if pinch is not None:
                last = pinch.lastScaleFactor() or 1.0
                self.zoom_by(pinch.scaleFactor() / last)
                event.accept()
                return True
        return super().event(event)

    def zoom_by(self, factor):
        factor = max(0.5, min(2.0, factor))
        self.viewer.setZoomMode(QPdfView.ZoomMode.Custom)
        self.viewer.setZoomFactor(max(0.15, min(8.0, self.viewer.zoomFactor() * factor)))

    def dragEnterEvent(self, event):
        urls = event.mimeData().urls() if event.mimeData().hasUrls() else []
        if any(url.isLocalFile() and Path(url.toLocalFile()).suffix.lower() == ".pdf" for url in urls):
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        pdfs = [Path(url.toLocalFile()) for url in event.mimeData().urls()
                if url.isLocalFile() and Path(url.toLocalFile()).suffix.lower() == ".pdf"]
        if not pdfs:
            event.ignore()
            return
        self.open_pdf(pdfs[0])
        event.acceptProposedAction()
        if len(pdfs) > 1:
            self.status.setText(f"Opened {pdfs[0].name}; drop one PDF at a time to choose another")

    def build_ui(self):
        toolbar = QToolBar("Document")
        toolbar.setMovable(False); self.addToolBar(toolbar)
        for label, callback, shortcut in [
            ("Open", self.choose_pdf, "Ctrl+O"), ("Save a copy", self.save_copy, "Ctrl+Shift+S"),
            ("−", self.zoom_out, "Ctrl+-"), ("+", self.zoom_in, "Ctrl++"), ("Fit width", self.fit_width, "Ctrl+0")]:
            action = QAction(label, self); action.triggered.connect(callback); action.setShortcut(shortcut); toolbar.addAction(action)
        toolbar.addSeparator(); toolbar.addWidget(self.status)

        side = QFrame(); side.setObjectName("sidebar"); box = QVBoxLayout(side); box.setContentsMargins(18, 18, 18, 18)
        title = QLabel("COMMAND PDF"); title.setObjectName("brand"); box.addWidget(title)
        box.addWidget(QLabel("PAGES")); box.addWidget(self.pages, 1)
        for heading, actions in [
            ("EDIT", [("Add text box", self.add_text), ("Rotate pages", self.rotate_pages), ("Delete pages", self.delete_pages)]),
            ("SIGN", [("Add visible signature", self.add_signature)]),
            ("CONVERT", [("Files to PDF", self.files_to_pdf), ("PDF to DOCX", self.pdf_to_docx), ("PDF to images", self.pdf_to_images), ("Extract text", self.extract_text), ("Merge PDFs", self.merge_pdfs)]),
        ]:
            label = QLabel(heading); label.setObjectName("section"); box.addWidget(label)
            for text, callback in actions:
                button = QPushButton(text); button.clicked.connect(callback); box.addWidget(button)
        splitter = QSplitter(); splitter.addWidget(side); splitter.addWidget(self.viewer); splitter.setSizes([220, 1060])
        self.setCentralWidget(splitter)
        self.setStyleSheet("""
            QMainWindow,QWidget{background:#08111d;color:#eaf2fb;font-family:'Noto Sans';font-size:13px}
            QToolBar{background:#0d1928;border-bottom:1px solid #22364d;padding:7px;spacing:5px}
            QToolButton,QPushButton{background:#14253a;color:#eaf2fb;border:1px solid #29425e;border-radius:6px;padding:8px 12px}
            QToolButton:hover,QPushButton:hover{background:#1b3551;border-color:#38bdf8}
            #sidebar{background:#0b1726;border-right:1px solid #22364d} #brand{font-size:20px;font-weight:800;color:#fff}
            #section{color:#48c8ff;font-size:11px;font-weight:800;margin-top:10px} #status{color:#9eb2c9;margin-left:8px}
            QListWidget,QPlainTextEdit,QLineEdit,QComboBox,QSpinBox{background:#07101b;border:1px solid #29425e;border-radius:5px;padding:5px;color:#eaf2fb}
            QListWidget::item{padding:7px} QListWidget::item:selected{background:#176fa4}
        """)

    def choose_pdf(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open PDF", str(Path.home()), "PDF documents (*.pdf)")
        if path: self.open_pdf(Path(path))

    def open_pdf(self, path):
        error = self.document.load(str(path))
        if error != QPdfDocument.Error.None_:
            QMessageBox.critical(self, "Cannot open PDF", f"Could not open:\n{path}")
            return
        self.current = path.resolve(); self.setWindowTitle(f"{path.name} — Command PDF"); self.status.setText(str(path))

    def document_ready(self, count):
        self.pages.clear()
        for index in range(count): self.pages.addItem(f"Page {index + 1}")
        if count: self.pages.setCurrentRow(0)

    def ensure_document(self):
        if self.current and self.current.exists(): return True
        QMessageBox.information(self, "No document", "Open a PDF first."); return False

    def go_page(self, page):
        if page >= 0: self.viewer.pageNavigator().jump(page, QPointF(), self.viewer.zoomFactor())

    def zoom_in(self): self.zoom_by(1.2)
    def zoom_out(self): self.zoom_by(1 / 1.2)
    def fit_width(self): self.viewer.setZoomMode(QPdfView.ZoomMode.FitToWidth)

    def save_copy(self):
        if not self.ensure_document(): return
        target = output_path(self, "Save a copy", self.current.with_stem(self.current.stem + "-copy"))
        if target: shutil.copy2(self.current, target); self.status.setText(f"Saved {target}")

    def edit_output(self, suffix):
        return output_path(self, "Save edited PDF", self.current.with_stem(self.current.stem + suffix))

    def position_rect(self, page_rect, choice, width, height):
        margin = 36
        x = margin if "left" in choice else page_rect.width - width - margin if "right" in choice else (page_rect.width-width)/2
        y = margin if "top" in choice else page_rect.height-height-margin if "bottom" in choice else (page_rect.height-height)/2
        return x, y, x + width, y + height

    def add_text(self):
        if not self.ensure_document(): return
        fitz = require_fitz(self)
        if not fitz: return
        dialog = TextDialog(self.document.pageCount(), self)
        if dialog.exec() != QDialog.Accepted or not dialog.text.toPlainText().strip(): return
        target = self.edit_output("-text-box")
        if not target: return
        doc = fitz.open(self.current); page = doc[dialog.page.value()-1]; choice = dialog.position.currentText().lower()
        width = min(dialog.width.value(), page.rect.width - 72)
        height = min(dialog.height.value(), page.rect.height - 72)
        rect = fitz.Rect(*self.position_rect(page.rect, choice, width, height))
        shape = page.new_shape()
        if dialog.background.isChecked():
            shape.draw_rect(rect); shape.finish(fill=(1, 1, 1), color=(0.72, 0.79, 0.86) if dialog.border.isChecked() else (1, 1, 1), width=1)
            shape.commit(overlay=True)
        elif dialog.border.isChecked():
            shape.draw_rect(rect); shape.finish(color=(0.40, 0.52, 0.64), width=1); shape.commit(overlay=True)
        inner = rect + (10, 8, -10, -8)
        remaining = page.insert_textbox(inner, dialog.text.toPlainText(), fontsize=dialog.size.value(), color=(0.05, 0.1, 0.16), lineheight=1.15)
        if remaining < 0:
            doc.close()
            QMessageBox.warning(self, "Text does not fit", "The text is too large for this box. Increase its height or reduce the font size.")
            return
        doc.save(target); doc.close(); self.open_pdf(target)

    def add_signature(self):
        if not self.ensure_document(): return
        fitz = require_fitz(self)
        if not fitz: return
        dialog = SignatureDialog(self.document.pageCount(), self)
        if dialog.exec() != QDialog.Accepted: return
        target = self.edit_output("-signed")
        if not target: return
        with tempfile.TemporaryDirectory(prefix="command-pdf-") as folder:
            signature = Path(folder) / "signature.png"; dialog.pad.image.save(str(signature), "PNG")
            doc = fitz.open(self.current); page = doc[dialog.page.value()-1]; width = dialog.width.value(); height = width * .36
            rect = self.position_rect(page.rect, dialog.position.currentText().lower(), width, height)
            page.insert_image(fitz.Rect(*rect), filename=str(signature), overlay=True)
            doc.save(target); doc.close()
        self.open_pdf(target)

    def rotate_pages(self):
        if not self.ensure_document(): return
        fitz = require_fitz(self)
        if not fitz: return
        text, ok = self.simple_prompt("Rotate pages", "Pages (example: 1,3-5; blank means all)", "")
        if not ok: return
        target = self.edit_output("-rotated")
        if not target: return
        doc = fitz.open(self.current)
        for index in self.parse_pages(text, len(doc)): doc[index].set_rotation((doc[index].rotation + 90) % 360)
        doc.save(target); doc.close(); self.open_pdf(target)

    def delete_pages(self):
        if not self.ensure_document(): return
        fitz = require_fitz(self)
        if not fitz: return
        text, ok = self.simple_prompt("Delete pages", "Pages to delete (example: 2,4-6)", "")
        if not ok or not text.strip(): return
        doc = fitz.open(self.current); pages = self.parse_pages(text, len(doc))
        if len(pages) >= len(doc): QMessageBox.warning(self, "Cannot delete", "A PDF must retain at least one page."); doc.close(); return
        target = self.edit_output("-pages-removed")
        if not target: doc.close(); return
        doc.delete_pages(pages); doc.save(target); doc.close(); self.open_pdf(target)

    def simple_prompt(self, title, label, value):
        dialog = QDialog(self); dialog.setWindowTitle(title); layout = QFormLayout(dialog); field = QLineEdit(value); layout.addRow(label, field)
        buttons = QDialogButtonBox(QDialogButtonBox.Cancel | QDialogButtonBox.Ok); buttons.accepted.connect(dialog.accept); buttons.rejected.connect(dialog.reject); layout.addRow(buttons)
        accepted = dialog.exec() == QDialog.Accepted
        return field.text(), accepted

    def parse_pages(self, text, count):
        if not text.strip(): return list(range(count))
        pages = set()
        try:
            for part in text.split(','):
                if '-' in part:
                    start, end = map(int, part.split('-', 1)); pages.update(range(start-1, end))
                else: pages.add(int(part)-1)
        except ValueError:
            QMessageBox.warning(self, "Invalid pages", "Use page numbers like 1,3-5."); return []
        return sorted(p for p in pages if 0 <= p < count)

    def files_to_pdf(self):
        files, _ = QFileDialog.getOpenFileNames(self, "Choose files", str(Path.home()), "Documents and images (*.png *.jpg *.jpeg *.webp *.txt *.odt *.doc *.docx *.ppt *.pptx *.xls *.xlsx)")
        if not files: return
        target = output_path(self, "Save converted PDF", Path(files[0]).with_suffix(".pdf"))
        if not target: return
        fitz = require_fitz(self)
        if not fitz: return
        converted = []
        with tempfile.TemporaryDirectory(prefix="command-pdf-convert-") as folder:
            folder = Path(folder)
            for item in map(Path, files):
                if item.suffix.lower() in {'.png','.jpg','.jpeg','.webp'}:
                    pdf = fitz.open(); page = pdf.new_page(); page.insert_image(page.rect, filename=str(item), keep_proportion=True); out = folder / f"{len(converted)}.pdf"; pdf.save(out); pdf.close(); converted.append(out)
                else:
                    run = subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(item)], capture_output=True, text=True)
                    out = folder / (item.stem + ".pdf")
                    if run.returncode == 0 and out.exists(): converted.append(out)
            if not converted: QMessageBox.critical(self, "Conversion failed", "None of the selected files could be converted."); return
            merged = fitz.open()
            for path in converted:
                source = fitz.open(path); merged.insert_pdf(source); source.close()
            merged.save(target); merged.close()
        self.open_pdf(target)

    def pdf_to_images(self):
        if not self.ensure_document(): return
        folder = QFileDialog.getExistingDirectory(self, "Choose image output folder", str(self.current.parent))
        if not folder: return
        prefix = str(Path(folder) / self.current.stem)
        result = subprocess.run(["pdftoppm", "-png", "-r", "150", str(self.current), prefix], capture_output=True, text=True)
        self.report_process(result, "Pages exported as PNG images.")

    def pdf_to_docx(self):
        if not self.ensure_document(): return
        try:
            from docx import Document
            from docx.enum.text import WD_BREAK
            from docx.shared import Inches, Pt
        except ImportError:
            QMessageBox.critical(self, "Conversion unavailable", "python-docx is missing. Run setup.sh in the Command PDF folder.")
            return
        fitz = require_fitz(self)
        if not fitz: return
        target, _ = QFileDialog.getSaveFileName(self, "Save Word document", str(self.current.with_suffix('.docx')), "Word documents (*.docx)")
        if not target: return
        source = fitz.open(self.current)
        document = Document()
        section = document.sections[0]
        if len(source):
            page_rect = source[0].rect
            section.page_width = Inches(page_rect.width / 72)
            section.page_height = Inches(page_rect.height / 72)
            section.top_margin = section.bottom_margin = Inches(.55)
            section.left_margin = section.right_margin = Inches(.65)
        for page_number, page in enumerate(source):
            blocks = sorted(page.get_text("blocks"), key=lambda block: (round(block[1] / 6), block[0]))
            for block in blocks:
                text = block[4].strip()
                if not text: continue
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                lines = text.splitlines()
                for index, line in enumerate(lines):
                    paragraph.add_run(line)
                    if index < len(lines) - 1: paragraph.add_run().add_break()
            if page_number < len(source) - 1:
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        source.close()
        document.save(target)
        self.status.setText(f"Word document saved to {target}")

    def extract_text(self):
        if not self.ensure_document(): return
        target, _ = QFileDialog.getSaveFileName(self, "Save extracted text", str(self.current.with_suffix('.txt')), "Text files (*.txt)")
        if not target: return
        result = subprocess.run(["pdftotext", "-layout", str(self.current), target], capture_output=True, text=True)
        self.report_process(result, f"Text saved to {target}")

    def merge_pdfs(self):
        files, _ = QFileDialog.getOpenFileNames(self, "Choose PDFs in merge order", str(Path.home()), "PDF documents (*.pdf)")
        if len(files) < 2: return
        target = output_path(self, "Save merged PDF", Path(files[0]).with_stem("merged-document"))
        if not target: return
        result = subprocess.run(["pdfunite", *files, str(target)], capture_output=True, text=True)
        if result.returncode == 0: self.open_pdf(target)
        else: self.report_process(result, "")

    def report_process(self, result, success):
        if result.returncode == 0: self.status.setText(success)
        else: QMessageBox.critical(self, "Operation failed", result.stderr.strip() or "The operation failed.")


def main():
    app = QApplication(sys.argv); app.setApplicationName("Command PDF"); app.setDesktopFileName("org.commandos.PDF")
    window = CommandPDF(sys.argv[1] if len(sys.argv) > 1 else None); window.show(); return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
